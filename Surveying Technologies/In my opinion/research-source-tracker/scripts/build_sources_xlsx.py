#!/usr/bin/env python3
"""Build or append to the research workbook (topic summaries + every source + search log).

Usage:
  build_sources_xlsx.py INPUT.json OUTPUT.xlsx

If OUTPUT.xlsx already exists, new rows are appended (nothing is overwritten, IDs continue
from the highest existing one, duplicate sources are skipped). Otherwise a new workbook is made.

INPUT.json (every top-level key is optional, but send at least one):
{
  "date": "2026-10-02",                       # default: today
  "summaries": [{
      "topic": "FIG commissions",
      "summary": "2-5 sentences, what we learned",
      "key_points": ["...", "..."],
      "open_questions": ["...", "..."],
      "source_refs": ["fig-comm", "fig-c5"]    # refs of sources below (or existing IDs like S003)
  }],
  "sources": [{
      "ref": "fig-comm",                        # local nickname, only used to link summaries
      "topic": "FIG commissions",
      "title": "FIG Commissions",
      "type": "Org website",                    # Web page | Journal paper | Book/PDF | Standard | Org website | Dataset | Other
      "author": "FIG",
      "year": "2025",
      "url": "https://www.fig.net/organisation/comm/index.asp",
      "doi": "",
      "location": "Commission 5 page",          # page number / section / chapter. For PDFs: 'p. 123 (PDF p. 145)'
      "key_info": "what was actually taken from this source",
      "reliability": "Verified - opened",       # see below
      "accessed": "2026-10-02",
      "search_keywords": "FIG commissions site:fig.net",
      "notes": ""
  }],
  "searches": [{
      "query": "FIG Commission 5 positioning working groups",
      "where": "Web search | Google Scholar | ResearchGate | Scopus | PDF search",
      "result": "what it returned / was it useful",
      "url": ""
  }]
}

Reliability values (anything else is stored as 'Unverified'):
  Verified - opened      the page / PDF was actually opened and the info read there
  Snippet only           seen only in a search-result snippet, not opened
  From memory - verify   recalled from training knowledge, no source opened (give search keys instead of links)
  Unverified             anything else
"""
import argparse
import datetime
import json
import math
import os
import re
import sys

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.formatting.rule import FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:
    sys.exit("openpyxl is not installed. Run: pip install openpyxl --break-system-packages")

# (header EN, header TH, width)
SHEETS = {
    "Summary": [
        ("Date", "วันที่", 12), ("Topic", "หัวข้อ", 26), ("Summary", "สรุป", 70),
        ("Key points", "ประเด็นสำคัญ", 58), ("Open questions / next steps", "ข้อสงสัย / ขั้นต่อไป", 44),
        ("Source IDs", "รหัสแหล่งที่มา", 16),
    ],
    "Sources": [
        ("ID", "รหัส", 8), ("Topic", "หัวข้อ", 24), ("Title", "ชื่อเรื่อง", 46), ("Type", "ประเภท", 14),
        ("Author / Org", "ผู้เขียน / องค์กร", 24), ("Year", "ปี", 7), ("Link", "ลิงก์", 44), ("DOI", "DOI", 22),
        ("Location (page / section)", "ตำแหน่ง (หน้า / หัวข้อ)", 24), ("Key info taken", "ข้อมูลที่นำมาใช้", 58),
        ("Reliability", "ความน่าเชื่อถือ", 21), ("Accessed", "วันที่เข้าถึง", 12),
        ("Search keywords", "คำค้น", 30), ("Notes", "หมายเหตุ", 34),
    ],
    "Search Log": [
        ("Date", "วันที่", 12), ("Query / keyword", "คำค้น", 46), ("Where searched", "ค้นที่ไหน", 22),
        ("Result / note", "ผลลัพธ์ / หมายเหตุ", 60), ("Link", "ลิงก์", 44),
    ],
}

RELIABILITY = ["Verified - opened", "Snippet only", "From memory - verify", "Unverified"]
REL_FILL = {"Verified - opened": "D9EAD3", "Snippet only": "FFF2CC",
            "From memory - verify": "FCE5CD", "Unverified": "E6E6E6"}

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(bold=True, color="FFFFFF")
THIN = Side(style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
LINK_FONT = Font(color="0563C1", underline="single")


def norm_reliability(v):
    s = (v or "").strip().lower()
    if s.startswith("verified"):
        return RELIABILITY[0]
    if s.startswith("snippet"):
        return RELIABILITY[1]
    if "memory" in s:
        return RELIABILITY[2]
    return RELIABILITY[3]


def as_text(v):
    if v is None:
        return ""
    if isinstance(v, (list, tuple)):
        items = [str(x).strip() for x in v if str(x).strip()]
        return "\n".join(f"• {x}" for x in items)
    return str(v).strip()


def set_cell(ws, row, col, value, link=False):
    c = ws.cell(row=row, column=col)
    c.value = value
    if isinstance(value, str) and value.startswith("="):
        c.data_type = "s"  # never let source text turn into a formula
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.border = BORDER
    if link and isinstance(value, str) and value.startswith(("http://", "https://")):
        c.hyperlink = value
        c.font = LINK_FONT
    return c


def ensure_sheet(wb, name):
    cols = SHEETS[name]
    if name in wb.sheetnames:
        return wb[name], False
    ws = wb.create_sheet(name)
    for j, (en, th, w) in enumerate(cols, 1):
        c = ws.cell(row=1, column=j, value=f"{en}\n{th}")
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[1].height = 34
    ws.freeze_panes = "A2"
    return ws, True


def fit_row_height(ws, row, ncols):
    lines = 1
    for j in range(1, ncols + 1):
        v = ws.cell(row=row, column=j).value
        if not isinstance(v, str):
            continue
        width = ws.column_dimensions[get_column_letter(j)].width or 12
        n = sum(max(1, math.ceil(len(part) / max(1, width * 1.05))) for part in v.split("\n"))
        lines = max(lines, n)
    ws.row_dimensions[row].height = min(15 * lines + 4, 330)


def last_data_row(ws):
    r = ws.max_row
    while r > 1 and all(ws.cell(row=r, column=c).value in (None, "") for c in range(1, ws.max_column + 1)):
        r -= 1
    return r


def key_of(url, title, location):
    base = (url or title or "").strip().lower().rstrip("/")
    return base + "|" + (location or "").strip().lower()


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("input")
    ap.add_argument("output")
    a = ap.parse_args()

    with open(a.input, encoding="utf-8") as f:
        data = json.load(f)
    if not any(data.get(k) for k in ("summaries", "sources", "searches")):
        sys.exit("input has no summaries, sources or searches - nothing to write")

    today = datetime.date.today().isoformat()
    day = data.get("date") or today
    warnings = []

    if os.path.exists(a.output):
        wb = load_workbook(a.output)
        mode = "appended to existing workbook"
    else:
        wb = Workbook()
        wb.remove(wb.active)
        mode = "created new workbook"
    ws_sum, _ = ensure_sheet(wb, "Summary")
    ws_src, _ = ensure_sheet(wb, "Sources")
    ws_log, _ = ensure_sheet(wb, "Search Log")
    # keep the sheet order stable even when appending to an older file
    wb._sheets = [wb[n] for n in SHEETS if n in wb.sheetnames] + [s for s in wb._sheets if s.title not in SHEETS]

    # ---- sources
    existing_ids, existing_keys = [], {}
    for r in range(2, last_data_row(ws_src) + 1):
        sid = ws_src.cell(row=r, column=1).value
        if not sid:
            continue
        existing_ids.append(int(m.group(1)) if (m := re.fullmatch(r"S(\d+)", str(sid))) else 0)
        existing_keys[key_of(ws_src.cell(row=r, column=7).value, ws_src.cell(row=r, column=3).value,
                             ws_src.cell(row=r, column=9).value)] = str(sid)
    next_id = (max(existing_ids) if existing_ids else 0) + 1
    ref_to_id, added_src, skipped_src = {}, 0, 0

    for s in data.get("sources", []):
        title, url, doi = as_text(s.get("title")), as_text(s.get("url")), as_text(s.get("doi"))
        if not url and doi:
            url = doi if doi.startswith("http") else f"https://doi.org/{doi}"
        if not (title or url):
            warnings.append("a source with neither title nor url was skipped")
            continue
        loc = as_text(s.get("location"))
        rel = norm_reliability(s.get("reliability"))
        k = key_of(url, title, loc)
        if k in existing_keys:
            if s.get("ref"):
                ref_to_id[s["ref"]] = existing_keys[k]
            skipped_src += 1
            continue
        sid = f"S{next_id:03d}"
        next_id += 1
        existing_keys[k] = sid
        if s.get("ref"):
            ref_to_id[s["ref"]] = sid
        if rel == RELIABILITY[0] and not (url or loc):
            warnings.append(f"{sid} '{title}' is marked Verified but has no link or page/section - add one")
        if rel != RELIABILITY[2] and not url and not loc:
            warnings.append(f"{sid} '{title}' has no link and no location - the user cannot trace it")
        row = last_data_row(ws_src) + 1
        vals = [sid, as_text(s.get("topic")), title, as_text(s.get("type")), as_text(s.get("author")),
                as_text(s.get("year")), url, doi, loc, as_text(s.get("key_info")), rel,
                as_text(s.get("accessed")) or day, as_text(s.get("search_keywords")), as_text(s.get("notes"))]
        for j, v in enumerate(vals, 1):
            set_cell(ws_src, row, j, v, link=(j == 7))
        fit_row_height(ws_src, row, len(vals))
        added_src += 1

    # ---- summaries (a row identical in date+topic+summary is skipped, so re-running after a failed save is safe)
    seen_sum = {(str(ws_sum.cell(row=r, column=1).value), str(ws_sum.cell(row=r, column=2).value),
                 str(ws_sum.cell(row=r, column=3).value))
                for r in range(2, last_data_row(ws_sum) + 1)}
    added_sum = 0
    for t in data.get("summaries", []):
        if (day, as_text(t.get("topic")), as_text(t.get("summary"))) in seen_sum:
            continue
        ids = []
        for ref in t.get("source_refs", []) or []:
            sid = ref_to_id.get(ref) or (ref if re.fullmatch(r"S\d+", str(ref)) else None)
            if sid:
                ids.append(sid)
            else:
                warnings.append(f"summary '{t.get('topic')}' refers to unknown source ref '{ref}'")
        row = last_data_row(ws_sum) + 1
        vals = [day, as_text(t.get("topic")), as_text(t.get("summary")), as_text(t.get("key_points")),
                as_text(t.get("open_questions")), ", ".join(dict.fromkeys(ids))]
        if not ids:
            warnings.append(f"summary '{t.get('topic')}' is not linked to any source")
        for j, v in enumerate(vals, 1):
            set_cell(ws_sum, row, j, v)
        fit_row_height(ws_sum, row, len(vals))
        added_sum += 1

    # ---- search log
    seen_log = {(str(ws_log.cell(row=r, column=1).value), str(ws_log.cell(row=r, column=2).value),
                 str(ws_log.cell(row=r, column=3).value))
                for r in range(2, last_data_row(ws_log) + 1)}
    added_log = 0
    for q in data.get("searches", []):
        if (day, as_text(q.get("query")), as_text(q.get("where"))) in seen_log:
            continue
        row = last_data_row(ws_log) + 1
        vals = [day, as_text(q.get("query")), as_text(q.get("where")), as_text(q.get("result")), as_text(q.get("url"))]
        for j, v in enumerate(vals, 1):
            set_cell(ws_log, row, j, v, link=(j == 5))
        fit_row_height(ws_log, row, len(vals))
        added_log += 1

    # ---- reliability dropdown + colours (re-applied so it covers appended rows)
    ws_src.data_validations.dataValidation = []
    dv = DataValidation(type="list", formula1='"' + ",".join(RELIABILITY) + '"', allow_blank=True,
                        showErrorMessage=False)
    ws_src.add_data_validation(dv)
    end = max(last_data_row(ws_src) + 200, 300)
    dv.add(f"K2:K{end}")
    ws_src.conditional_formatting = type(ws_src.conditional_formatting)()
    for name, color in REL_FILL.items():
        ws_src.conditional_formatting.add(
            f"K2:K{end}", FormulaRule(formula=[f'$K2="{name}"'], fill=PatternFill("solid", bgColor=color, fgColor=color)))
    for ws in (ws_sum, ws_src, ws_log):
        ws.page_setup.orientation = "landscape"  # prints/exports on one page width
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.auto_filter.ref = f"A1:{get_column_letter(len(SHEETS[ws.title]))}{max(last_data_row(ws), 1)}"

    wb.active = 0
    wb.save(a.output)

    # ---- verify by reloading
    chk = load_workbook(a.output)
    print(f"{mode}: {a.output}")
    print(f"  added: {added_sum} summary row(s), {added_src} source(s), {added_log} search log row(s); "
          f"duplicates skipped: {skipped_src}")
    print("  sheet sizes now: " + ", ".join(f"{n}={last_data_row(chk[n]) - 1} rows" for n in SHEETS))
    for w in warnings:
        print(f"  WARNING: {w}")


if __name__ == "__main__":
    main()
